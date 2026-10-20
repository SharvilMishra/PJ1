"""Normalize the supplied workbook into the phone catalog used by web/API."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Datasets" / "sic_mobile_spec_PJ1.xlsx"
OUTPUT = ROOT / "Frontend" / "src" / "data" / "phones.json"


def number(value: object, *, storage: bool = False) -> float | None:
    if value is None or str(value).strip().lower() in {"", "not mentioned", "n/a", "na", "-"}:
        return None
    text = str(value).strip().replace(",", "")
    multiplier = 1024 if storage and re.search(r"(tb|terabyte)s?\s*$", text, re.I) else 1
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) * multiplier if match else None


def slug(value: str) -> str:
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", value.lower())).strip("-")


def percentile(values: list[float], fraction: float) -> float:
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source workbook: {SOURCE}")
    sheet = load_workbook(SOURCE, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip().lower() for value in next(rows)]
    source_rows = [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]

    prices = [number(row.get("price")) for row in source_rows]
    prices = [value for value in prices if value is not None]
    cameras = [number(row.get("camera")) for row in source_rows]
    cameras = [value for value in cameras if value is not None and value > 0]
    price_budget = percentile(prices, 0.33)
    price_flagship = percentile(prices, 0.75)
    camera_high = percentile(cameras, 0.75)

    phones = []
    seen_ids: set[str] = set()
    for row in source_rows:
        name = str(row.get("name") or "").strip()
        if not name:
            continue
        phone_id = slug(name)
        if phone_id in seen_ids:
            raise SystemExit(f"Phone names must be unique; duplicate: {name}")
        seen_ids.add(phone_id)

        ram = number(row.get("ram"))
        price = number(row.get("price"))
        camera = number(row.get("camera"))
        battery = number(row.get("battery"))
        categories = []
        if ram is not None and ram >= 8:
            categories.append("gaming")
        if battery is not None and battery >= 5000:
            categories.append("battery")
        if camera is not None and camera > 0 and camera >= camera_high:
            categories.append("camera")
        if price is not None and price <= price_budget:
            categories.append("budget")
        if (price is not None and price >= price_flagship) or (ram is not None and ram >= 12):
            categories.append("flagship")

        parts = name.split(maxsplit=1)
        processor = str(row.get("processor") or "").strip()
        phones.append({
            "id": phone_id,
            "name": name,
            "brand": parts[0],
            "model": parts[1] if len(parts) > 1 else parts[0],
            "rating": number(row.get("rating")),
            "price": price,
            "categories": categories,
            "cameraMainMp": camera if camera and camera > 0 else None,
            "displayType": str(row.get("display_type") or "").strip() or None,
            "displaySizeIn": number(row.get("display_size")),
            "batteryMah": battery,
            "storageGb": number(row.get("storage"), storage=True),
            "ramGb": ram,
            "weightG": number(row.get("weight")),
            "processor": None if processor.lower() == "not mentioned" else processor or None,
            "os": None,
            "imageUrl": None,
        })

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(phones, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(phones)} phones to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
